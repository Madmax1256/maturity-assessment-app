"""Genera casos de paridad entre el motor y el Excel V01.

Para cada caso escribe respuestas y alcance en una copia del Excel, la recalcula con
LibreOffice en modo headless y guarda lo que muestran las fórmulas del libro
(INICIO E36:E47, X36:X47, J36 y M39) junto con las entradas, en fixtures/parity.json.

Uso: python3 tools/parity/make_cases.py [ruta.xlsx] [n_completos] [n_parciales]
Requiere openpyxl y LibreOffice (soffice).
"""
import json
import random
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "fixtures" / "Diagnostico_FS_V01.xlsx"
N_FULL = int(sys.argv[2]) if len(sys.argv) > 2 else 20
N_PART = int(sys.argv[3]) if len(sys.argv) > 3 else 6
WORK = ROOT / "fixtures" / "parity-work"
OUT = ROOT / "fixtures" / "parity.json"
N_Q = 128
FIRST_ROW = 5  # Preguntas!F5 = pregunta 1


def build_cases():
    rnd = random.Random(20261008)
    cases = []
    for i in range(N_FULL + N_PART):
        partial = i >= N_FULL
        style = rnd.choice(["uniforme", "alto", "bajo", "extremos"])
        weights = {"uniforme": [1, 1, 1, 1, 1], "alto": [1, 1, 2, 4, 3], "bajo": [3, 4, 2, 1, 1], "extremos": [4, 1, 0, 1, 4]}[style]
        answers = [rnd.choices(range(5), weights)[0] for _ in range(N_Q)]
        if i == 0:
            answers = [4] * N_Q  # todo 100 %
        if i == 1:
            answers = [0] * N_Q  # todo 0 %
        if partial:
            for k in rnd.sample(range(N_Q), rnd.randint(1, 40)):
                answers[k] = None
        applies = [True] * 12
        if i >= 2 and rnd.random() < 0.5:
            for d in rnd.sample(range(12), rnd.randint(1, 4)):
                applies[d] = False
        cases.append({"id": f"caso-{i + 1:02d}", "style": style, "partial": partial, "answers": answers, "applies": applies})
    return cases


def write_inputs(cases):
    if WORK.exists():
        shutil.rmtree(WORK)
    (WORK / "in").mkdir(parents=True)
    for c in cases:
        wb = openpyxl.load_workbook(SRC)
        p, lists, ini = wb["Preguntas"], wb["Listas Evaluación"], wb["INICIO - DASHBOARD"]
        for k, s in enumerate(c["answers"]):
            p[f"F{FIRST_ROW + k}"] = "EVA" if s is None else lists.cell(row=k + 1, column=2 + s).value
        for d, ok in enumerate(c["applies"]):
            ini[f"D{36 + d}"] = "SI" if ok else "NO"
        wb.save(WORK / "in" / f"{c['id']}.xlsx")


def recalc():
    files = sorted(str(f) for f in (WORK / "in").glob("*.xlsx"))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(WORK / "out"), *files],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)


def read_outputs(cases):
    for c in cases:
        wb = openpyxl.load_workbook(WORK / "out" / f"{c['id']}.xlsx", data_only=True)
        ini = wb["INICIO - DASHBOARD"]
        c["excel"] = {
            "dimensions": [ini[f"E{r}"].value for r in range(36, 48)],
            "dimensionLevels": [ini[f"X{r}"].value for r in range(36, 48)],
            "global": ini["J36"].value,
            "globalLabel": ini["M39"].value,
        }


def main():
    cases = build_cases()
    write_inputs(cases)
    recalc()
    read_outputs(cases)
    OUT.write_text(json.dumps({"source": SRC.name, "cases": cases}, ensure_ascii=False, indent=1) + "\n")
    print(f"{len(cases)} casos escritos en {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
